from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


ROOT = Path(r"C:\Users\Administrator\Desktop\project\pre-diagnosis\intake-diagnostician\projects\project_report_regular_ppt169_20260918")
RENDERED = ROOT / "rendered_v2"
OUTPUT = ROOT / "exports" / "project_report_regular_compatible_v2.pptx"


prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)

for index in range(1, 15):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    image = RENDERED / f"slide-{index}.png"
    slide.shapes.add_picture(
        str(image),
        0,
        0,
        width=prs.slide_width,
        height=prs.slide_height,
    )

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(OUTPUT)
