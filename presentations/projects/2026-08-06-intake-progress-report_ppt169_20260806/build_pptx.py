"""
Build PPTX from SVG slides for intake-diagnostician progress report.
Converts each SVG to PNG and embeds as a full-slide image.
"""
import io
import os
import sys
from pathlib import Path

import cairosvg
from pptx import Presentation
from pptx.util import Inches, Pt, Emu

# ── Paths ──
PROJECT_DIR = Path(__file__).resolve().parent
SVG_DIR = PROJECT_DIR / "svg_output"
OUTPUT_DIR = Path(r"C:\Users\Administrator\Desktop\codex\intake-diagnostician\ppt-output")
OUTPUT_FILE = OUTPUT_DIR / "中医诊前问诊系统-进展汇报-2026-08-06.pptx"

# Slide dimensions: 16:9 at 1280x720 → 13.333" x 7.5"
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# ── SVG files in order ──
SVG_FILES = [
    "01-cover.svg",
    "02-feedback.svg",
    "03-workflow.svg",
    "04-skills.svg",
    "05-followup-loop.svg",
    "06-analysis-scope.svg",
    "07-engineering.svg",
    "08-conclusion.svg",
]

# ── Speaker notes per slide ──
NOTES = {
    "01-cover.svg": (
        "各位老师和同学大家好，今天我汇报的是中医诊前问诊系统的最新进展。"
        "本轮迭代的核心是三个关键词：完整度判断、动态补问和检查建议。"
        "我们把系统从一个固定问卷工具，升级为能够动态识别信息缺口并主动追问的信息采集系统。"
    ),
    "02-feedback.svg": (
        "首先回顾上一次组会的反馈。"
        "组会指出，完整度不能只判断十道固定题目是否全部回答，这太机械了。"
        "应该结合十问歌框架和中医问诊经验，判断当前采集的信息是否能支撑后续诊中使用。"
        "同时，辨证参考在诊前阶段风险太高，决定删除，只保留信息完整度和推荐检查项目。"
        "另外，将原本混在一起的能力拆分为两个职责清晰的Skill。"
    ),
    "03-workflow.svg": (
        "在新的流程设计中，患者首先完成基础问诊，采集主诉和十问歌项目。"
        "然后由完整度Skill判断当前信息的缺口。如果有缺失的必问项，就进入动态补问环节。"
        "补问环节由追问Skill负责生成问题，用户回答后，不是直接结束，而是重新交回完整度Skill评估。"
        "这个「回答后重新评估」的闭环是本次流程升级的核心。"
        "最终产出的信息档案和检查建议可以直接供后续诊中环节参考。"
    ),
    "04-skills.svg": (
        "两个Skill的分工非常明确。"
        "Skill 1，tcm-intake-checklist，回答'应该问什么'。它依据十问歌框架、基础安全项和老中医问诊经验，"
        "输出缺失的必问项、可选补充项和当前的完整度状态。"
        "Skill 2，tcm-questioning-guide，回答'应该怎么问'。它负责控制提问的语气和顺序，"
        "一次只问一个关键问题，使用通俗自然不诱导的表达，不重复已回答的内容，不做诊断，不暗示处方。"
        "两个Skill的内容合并注入一次千问模型调用，完整度判断的输出直接成为追问生成的输入。"
    ),
    "05-followup-loop.svg": (
        "这张图展示了动态补问的完整闭环。"
        "固定问答完成不等于信息已经完整，这是核心认识。"
        "流程从完整度Skill判断开始，识别当前最关键的信息缺口，然后由追问Skill生成一个问题。"
        "用户回答后，回答记录写入信息档案，然后重新交回完整度Skill判断。"
        "我们设定了最多5轮动态补问的上限，避免过度追问。"
        "当完整度Skill判断不需要继续追问或者达到5轮上限时，闭环结束，产出最终的信息档案和完整度报告。"
    ),
    "06-analysis-scope.svg": (
        "在分析输出方面，我们做了明确的收敛。"
        "保留的内容包括：信息完整度状态、缺失必问项、可选补充项、安全提示，以及推荐检查项目。"
        "检查项目包含六个标准化字段：项目名称、检查类型、推荐科室、临床意义、推荐优先级和检查前注意事项。"
        "删除的内容包括：辨证参考、疾病诊断、处方建议和检查价格。"
        "删除这些内容的根本原因是：诊前问诊阶段信息不完整，无法做出可靠的辨证诊断，"
        "诊断和处方应统一留给诊中环节。同时每份分析结果末尾都有免责声明。"
    ),
    "07-engineering.svg": (
        "在工程实现方面，我们做了几个关键的稳定性设计。"
        "第一，分析Agent与普通问诊Agent完全分离，各自使用独立的thread_id隔离上下文。"
        "第二，两个Skill的内容合并注入一次千问模型调用，避免多次API调用增加耗时。"
        "第三，使用后台线程执行智能分析，不阻塞前端交互。"
        "第四，设置45秒超时，超时后返回可重试的结果，前端提供重新分析按钮。"
        "第五，结构化保存三部分数据：基础回答、动态补问和分析结果。"
        "验证结果显示，后端22项测试全部通过，前端生产构建通过，前后端健康检查正常。"
        "另外还清理了前端右侧冗余的基础问项列表。"
    ),
    "08-conclusion.svg": (
        "最后总结一下。本轮我们完成了六项核心改造：完整度判断逻辑重构，两个Skill职责拆分，"
        "动态补问闭环，智能分析范围收敛，检查建议结构化，以及稳定性优化与测试。"
        "下一步的工作重点包括：继续细化必问项、选问项和判断规则；补充年龄、过敏史等采集字段；"
        "由临床人员审核检查推荐规则；删除遗留未使用的辨证代码；使用更多真实问诊场景测试。"
        "总的来说，系统已经从固定问卷升级为可动态补齐关键缺口的诊前信息工具。"
        "谢谢大家，欢迎提问。"
    ),
}

os.makedirs(OUTPUT_DIR, exist_ok=True)


def svg_to_png_bytes(svg_path: Path) -> io.BytesIO:
    """Convert SVG file to PNG bytes using cairosvg."""
    svg_content = svg_path.read_text(encoding="utf-8")
    png_bytes = cairosvg.svg2png(
        bytestring=svg_content.encode("utf-8"),
        output_width=1280,
        output_height=720,
    )
    buf = io.BytesIO(png_bytes)
    return buf


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # Use blank layout
    blank_layout = prs.slide_layouts[6]  # blank

    for svg_file in SVG_FILES:
        svg_path = SVG_DIR / svg_file
        if not svg_path.exists():
            print(f"ERROR: {svg_path} not found!")
            sys.exit(1)

        print(f"Processing {svg_file}...")

        # Convert SVG to PNG
        try:
            png_buf = svg_to_png_bytes(svg_path)
        except Exception as e:
            print(f"ERROR converting {svg_file}: {e}")
            sys.exit(1)

        # Add slide
        slide = prs.slides.add_slide(blank_layout)

        # Add PNG as full-slide image
        slide.shapes.add_picture(
            png_buf,
            Inches(0),
            Inches(0),
            SLIDE_W,
            SLIDE_H,
        )

        # Add speaker notes
        note_text = NOTES.get(svg_file, "")
        if note_text:
            notes_slide = slide.notes_slide
            notes_slide.notes_text_frame.text = note_text

    # Save PPTX
    print(f"\nSaving to {OUTPUT_FILE}...")
    prs.save(str(OUTPUT_FILE))
    print(f"Done! PPTX saved to: {OUTPUT_FILE}")
    print(f"Total slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
