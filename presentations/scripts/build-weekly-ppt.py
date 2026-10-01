"""
生成本周进展 PPT — 7 页
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ── 颜色常量 ──
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
DARK   = RGBColor(0x1E, 0x29, 0x3B)
GRAY   = RGBColor(0x64, 0x74, 0x8B)
LIGHT  = RGBColor(0xE2, 0xE8, 0xF0)
ACCENT = RGBColor(0x4F, 0x46, 0xE5)  # indigo-600
TAG_BG = RGBColor(0xEE, 0xF2, 0xFF)  # indigo-50
CARD_BG = RGBColor(0xF8, 0xFA, 0xFC)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)

# ── 辅助函数 ──

def add_text(slide, left, top, width, height, text, size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return txBox

def add_rect(slide, left, top, width, height, fill=CARD_BG, border=LIGHT, radius=None):
    shape = slide.shapes.add_shape(
        5 if radius else 1,  # 5=roundRect, 1=rect
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(1)
    return shape

def add_divider(slide, top):
    add_rect(slide, 0.75, top, 11.83, 0.01, fill=LIGHT, border=LIGHT)

def add_title(slide, text):
    add_text(slide, 0.75, 0.5, 11.83, 0.6, text, size=34, bold=True)
    add_divider(slide, 1.15)

def add_card(slide, x, y, w, h, title_text, body_text):
    """画一张带标题的卡片"""
    add_rect(slide, x, y, w, h)
    if title_text:
        add_text(slide, x + 0.2, y + 0.15, w - 0.4, 0.35, title_text, size=15, bold=True, color=ACCENT)
    add_text(slide, x + 0.2, y + (0.55 if title_text else 0.15), w - 0.4, h - (0.7 if title_text else 0.3), body_text, size=14, color=DARK)

def add_tag(slide, x, y, text):
    shape = add_rect(slide, x, y, 1.0, 0.38, fill=TAG_BG, border=LIGHT, radius=True)
    shape.text_frame.word_wrap = True
    p = shape.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(12)
    p.font.color.rgb = ACCENT
    p.alignment = PP_ALIGN.CENTER

# ================================================================
# SLIDE 1: 封面
# ================================================================
s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
bg = s.background
bg.fill.solid()
bg.fill.fore_color.rgb = WHITE

add_text(s, 0.75, 2.2, 11.83, 1.0, "AI 老中医问诊 · 本周进展", size=48, bold=True)
add_text(s, 0.75, 3.4, 11.83, 0.6, "前端风格焕新  ·  中医知识 Skill 体系", size=26, color=GRAY)
add_rect(s, 0.75, 4.3, 1.2, 0.04, fill=ACCENT, border=ACCENT)
add_text(s, 0.75, 4.8, 11.83, 0.5, "2026.07.22  ·  Deep Agents v4.0", size=18, color=RGBColor(0x94, 0xA3, 0xB8))

# ================================================================
# SLIDE 2: 本周概览
# ================================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
bg = s.background; bg.fill.solid(); bg.fill.fore_color.rgb = WHITE
add_title(s, "本周概览")

rows = [
    ("维度", "进展"),
    ("🎨 前端", "整体视觉风格升级：暖调古风 → 现代中式融合"),
    ("🧠 技能", "新增 2 个中医 Skill：tcm-intake（问诊指南）+ tcm-zhengxing（辨证速查）"),
    ("🏗️ 架构", "Deep Agents 框架稳定运行，4 个子智能体按需并行辨证"),
]
for i, (label, content) in enumerate(rows):
    y = 1.7 + i * 0.75
    add_text(s, 0.75, y, 2.5, 0.5, label, size=18, bold=(i == 0), color=DARK if i > 0 else ACCENT)
    add_text(s, 3.25, y, 9.33, 0.5, content, size=18, bold=(i == 0), color=DARK)
    if i == 0:
        add_rect(s, 0.75, y + 0.5, 11.83, 0.01, fill=LIGHT, border=LIGHT)

# ================================================================
# SLIDE 3: 前端风格对比
# ================================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
bg = s.background; bg.fill.solid(); bg.fill.fore_color.rgb = WHITE
add_title(s, "前端 · 风格对比")

add_card(s, 0.75, 1.6, 5.5, 3.8, "旧风格",
    "▸ 宣纸黄背景 (#F5F0E8)\n"
    "▸ 居中单栏聊天\n"
    "▸ 暖黄色聊天气泡\n"
    "▸ 翠绿强调色 (#2E7D32)\n"
    "▸ 全楷体字体\n"
    "▸ 无侧边导航\n\n"
    "Before（旧版）")

add_card(s, 7.08, 1.6, 5.5, 3.8, "新风格",
    "▸ 薰衣草蓝白渐变背景\n"
    "▸ 260px 左侧导航栏 + 主聊天区\n"
    "▸ 白卡片 + 蓝灰边框气泡\n"
    "▸ 薰衣草蓝强调色 (#6B7DB3)\n"
    "▸ 标题楷体 + 正文系统无衬线\n"
    "▸ 侧边栏功能导航 + 十问歌速览\n\n"
    "After（新版）")

add_text(s, 5.9, 3.2, 1.5, 0.8, "→", size=40, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
add_text(s, 0.75, 5.8, 11.83, 0.8,
    "配色思路：保留楷体标题的中医辨识度，用薰衣草蓝 + 白色卡片营造干净现代的质感，侧边栏提升功能可见性",
    size=15, color=GRAY)

# ================================================================
# SLIDE 4: 前端布局
# ================================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
bg = s.background; bg.fill.solid(); bg.fill.fore_color.rgb = WHITE
add_title(s, "前端 · 新布局结构")

add_rect(s, 1.0, 1.7, 2.6, 5.0, fill=TAG_BG, border=LIGHT, radius=True)
add_text(s, 1.15, 1.85, 2.3, 4.7,
    "侧边栏\n\n🍂 AI老中医\n  中医诊前助手\n\n──────────\n💬 开始问诊\n📋 问诊档案\nℹ️ 关于系统\n\n──────────\n十问歌\n一问寒热二问汗\n三问头身四问便\n五问饮食六问胸\n七聋八渴俱当辨\n\n──────────\n🟢 Deep Agents\n    v4.0",
    size=12, color=DARK)

# Top bar
add_rect(s, 3.9, 1.7, 8.68, 0.7, fill=WHITE, border=LIGHT, radius=True)
add_text(s, 4.1, 1.78, 8.0, 0.55, "顶栏  🍂 小郎中问诊  [诊前采集]                              🕐 实时问诊中", size=13, color=DARK)

# Chat area
add_rect(s, 3.9, 2.6, 8.68, 3.1, fill=CARD_BG, border=LIGHT, radius=True)
add_text(s, 4.1, 2.8, 5.0, 1.0, "🏺 助手消息\n白卡片 + 蓝灰细边框\n圆角左下小角标", size=13, color=GRAY)
add_text(s, 8.0, 3.8, 4.0, 0.6, "🌿 用户消息\n蓝紫渐变底 + 白色字  →", size=13, color=GRAY, align=PP_ALIGN.RIGHT)

# Input
add_rect(s, 3.9, 5.9, 8.68, 0.7, fill=WHITE, border=LIGHT, radius=True)
add_text(s, 4.1, 6.0, 7.0, 0.5, "✍️ 输入区  ·  蓝紫发送按钮", size=13, color=GRAY)

add_text(s, 0.75, 6.9, 11.83, 0.4,
    "TailwindCSS + 内联样式 · Vue 3 Composition API · markdown-it 表格渲染 · fetch + Vite proxy 前后端通信",
    size=13, color=GRAY)

# ================================================================
# SLIDE 5: Skill ①
# ================================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
bg = s.background; bg.fill.solid(); bg.fill.fore_color.rgb = WHITE
add_title(s, "Skill ①  十问歌深度问诊指南")
add_text(s, 0.75, 1.25, 6.0, 0.35, "skills/tcm-intake/SKILL.md", size=13, color=RGBColor(0x94, 0xA3, 0xB8))

# Tags
tags = ["寒热", "汗出", "头身", "二便", "饮食", "胸腹", "睡眠", "口渴", "旧病", "病因", "经期", "舌象"]
tx, ty = 0.75, 1.75
for i, tag in enumerate(tags):
    add_tag(s, tx, ty, tag)
    tx += 1.08
    if tx > 7.0:
        tx = 0.75
        ty += 0.5

add_card(s, 0.75, 3.0, 5.5, 2.8, "每个主题包含",
    "主问题：一句温和的提问\n\n"
    "追问方向：3-6 个细分追问路径\n"
    "  例：头痛 → 追问痛感 / 位置 / 持续时间 / 诱因\n\n"
    "记录要点：该项需记录的关键信息\n\n"
    "红旗信号（6 种需建议立即就医）\n"
    "  剧烈头痛伴呕吐 | 胸痛彻背 | 突然半身麻木\n"
    "  高热不退 | 剧烈腹痛拒按 | 呕血便血")

add_card(s, 7.08, 3.0, 5.5, 2.8, "效果对比",
    "Before（无 Skill）\n"
    "\"好的，头疼我记下了。接下来你出汗多吗？\"\n"
    "→ 跳下一题，无追问\n\n"
    "After（有 Skill 自动查阅）\n"
    "\"这个头痛，是胀痛还是刺痛？两边太阳穴还是\n"
    "后脑勺？每天什么时候最明显？有没有想吐？\"\n"
    "→ 专业追问，信息密度大幅提升")

add_text(s, 0.75, 6.2, 11.83, 0.5,
    "12 个追问维度  ·  60+ 细分追问方向  ·  6 种红旗信号  ·  按需加载，不占用启动上下文",
    size=15, color=GRAY)

# ================================================================
# SLIDE 6: Skill ②
# ================================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
bg = s.background; bg.fill.solid(); bg.fill.fore_color.rgb = WHITE
add_title(s, "Skill ②  辨证速查表")
add_text(s, 0.75, 1.25, 8.0, 0.35, "skills/tcm-zhengxing/SKILL.md  ·  供 4 个子智能体并行分析时查阅", size=13, color=RGBColor(0x94, 0xA3, 0xB8))

systems = [
    ("八纲辨证", "10+ 证型", "表里 · 寒热 · 虚实 · 阴阳", 0.75, 1.85),
    ("脏腑辨证", "30+ 证型", "心系 · 肝系 · 脾系 · 肺系 · 肾系", 4.5, 1.85),
    ("气血津液", "11 种病证", "气虚 · 气滞 · 血瘀 · 痰证 · 水停", 8.25, 1.85),
    ("病因病机", "六淫 + 七情", "风 · 寒 · 暑 · 湿 · 燥 · 火\n怒 · 喜 · 忧 · 思 · 悲 · 恐 · 惊", 0.75, 4.3),
    ("复合证型", "7 种组合", "肝郁脾虚 · 心肾不交 · 脾肾阳虚 · 气阴两虚\n痰瘀互结 · 肝阳上亢 · 上热下寒", 4.5, 4.3),
]

for name, count, items, x, y in systems:
    add_card(s, x, y, 3.45, 2.15, name, f"{count}\n\n{items}")

# Usage flow
add_card(s, 8.25, 4.3, 4.33, 2.15, "使用方式",
    "子智能体分析流程\n\n"
    "1. 八纲辨证 → 定大方向\n"
    "2. 脏腑辨证 → 定位病变脏腑\n"
    "3. 气血津液 → 分析物质基础\n"
    "4. 病因病机 → 推断病因\n"
    "→ 综合输出：病位+病性+核心病机+证型")

add_text(s, 0.75, 6.7, 11.83, 0.5,
    "覆盖 50+ 常见证型  ·  每证型含主症 + 关键鉴别点  ·  子智能体分析时有据可查",
    size=15, color=GRAY)

# ================================================================
# SLIDE 7: 总结
# ================================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
bg = s.background; bg.fill.solid(); bg.fill.fore_color.rgb = WHITE
add_title(s, "总结")

dones = [
    "前端全新视觉风格 — 薰衣草蓝白 + 侧边栏 + 现代卡片",
    "tcm-intake Skill — 十二维度深度追问指南 + 红旗信号",
    "tcm-zhengxing Skill — 四大辨证体系速查表 50+ 证型",
    "Deep Agents 框架稳定运行 — 4 子智能体并行辨证",
    "Skills 按需加载 — 不触发不加载，不占启动上下文",
]
for i, d in enumerate(dones):
    add_text(s, 0.75, 1.8 + i * 0.55, 11.83, 0.45, f"✅  {d}", size=17, color=DARK)

add_divider(s, 4.7)

add_text(s, 0.75, 5.0, 11.83, 0.5, "下一步", size=22, bold=True, color=ACCENT)

nexts = [
    "药膳食疗 Skill — 基于证型推荐食疗方案",
    "舌象识别 — 拍照自动分析舌色/苔色/苔质",
    "问诊档案检索与体质追踪",
]
for i, n in enumerate(nexts):
    add_text(s, 0.75, 5.6 + i * 0.45, 11.83, 0.4, f"🔜  {n}", size=16, color=GRAY)

# ── 导出 ──
out = "ppt-output/AI老中医问诊-本周进展.pptx"
prs.save(out)
print(f"Done → {out}")
